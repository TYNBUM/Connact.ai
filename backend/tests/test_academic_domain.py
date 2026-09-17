from io import BytesIO

import pytest
from docx import Document

from app.config import settings
from app.db import Session, WorkspaceRepository
from app.models import Contact, PersonaRevision
from app.services.contacts import contact_json, upsert_search
from app.services.documents import process_document, stop_document_worker
from app.services.people_jobs import process_job, stop_people_worker


@pytest.fixture
def paused_workers(client):
    stop_people_worker()
    stop_document_worker()
    yield client
    stop_people_worker()
    stop_document_worker()


def academic_persona(client, **values):
    payload = {
        "domain": "academic",
        "label": "2027 research outreach",
        "data": {
            "name": "Alex Researcher",
            "education": "Example University, MSc",
            "experience": "Research assistant in medical imaging",
            "skills": "Python, deep learning",
            "sectors": "Machine learning and medical imaging",
            "career_goals": "Pursue doctoral research",
            "target_regions": "United States",
            "target_roles": "Redwood University",
            "contact_purpose": "Ask about research fit",
        },
        **values,
    }
    response = client.post("/api/personas", json=payload)
    assert response.status_code == 200, response.text
    return response.json()


def manual_contact(client, domain, name):
    response = client.post(
        "/api/contacts",
        json={
            "domain": domain,
            "name": name,
            "title": "Professor",
            "company": "Example University",
            "sector": "Medical imaging" if domain == "academic" else "Banking",
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_academic_search_jobs_history_assessment_and_finance_isolation(
    paused_workers,
):
    client = paused_workers
    sync_search = client.post(
        "/api/academic/search", json={"title": " Professor "}
    )
    assert sync_search.status_code == 200
    assert sync_search.json()["total"] == 12
    finance = client.post("/api/finance/search/jobs", json={}).json()["job"]
    academic = client.post("/api/academic/search/jobs", json={}).json()["job"]
    assert finance["domain"] == "finance"
    assert academic["domain"] == "academic"
    assert finance["id"] != academic["id"]
    normalized = client.post(
        "/api/academic/search/jobs", json={"title": " Professor "}
    ).json()
    assert normalized["cached"] and normalized["job"]["id"] == academic["id"]

    assert process_job(finance["id"]) is None
    assert process_job(academic["id"]) is None
    finance_done = client.get("/api/people/jobs/" + finance["id"]).json()
    academic_done = client.get("/api/people/jobs/" + academic["id"]).json()
    assert finance_done["status"] == academic_done["status"] == "succeeded"
    assert academic_done["result"]["total"] == 12
    assert academic_done["result"]["domain"] == "academic"
    assert all(not item["email"] for item in academic_done["result"]["items"])
    mentor = academic_done["result"]["items"][0]
    assert mentor["domains"]["academic"]["sector"]
    assert "finance" not in mentor["domains"]
    assert "academic" not in finance_done["result"]["items"][0]["domains"]

    assert [job["id"] for job in client.get("/api/academic/search/jobs").json()] == [
        academic["id"]
    ]
    assert [job["id"] for job in client.get("/api/finance/search/jobs").json()] == [
        finance["id"]
    ]
    cached_academic = client.post("/api/academic/search/jobs", json={}).json()
    assert cached_academic["cached"]
    assert cached_academic["job"]["id"] == academic["id"]

    persona = academic_persona(client)
    assessment = client.post(
        "/api/academic/assess",
        json={"contact_ids": [mentor["id"]], "persona_id": persona["id"]},
    )
    assert assessment.status_code == 200, assessment.text
    assert "research-fit" in assessment.json()[0]["reason"]
    assert (
        client.post(
            "/api/finance/assess",
            json={"contact_ids": [mentor["id"]], "persona_id": persona["id"]},
        ).status_code
        == 422
    )


def test_same_linkedin_identity_gets_two_profiles_without_erasing_research(client):
    base = {
        "provider": "serpapi",
        "provider_id": "same-linkedin-person",
        "name": "Jamie Scholar",
        "title": "Professor",
        "company": "Example University",
        "location": "Boston, US",
        "school": "Example University",
        "profile_url": "https://www.linkedin.com/in/jamie-scholar",
        "email": "",
        "email_status": "not_requested",
    }
    with Session() as db:
        repo = WorkspaceRepository(db, settings.workspace_id)
        finance = upsert_search(repo, {**base, "sector": "Asset Management"})
        academic = upsert_search(repo, {**base, "sector": "Medical Imaging"}, "academic")
        assert academic.id == finance.id
        # A later live SerpAPI discovery has no verified research field and must not clear it.
        again = upsert_search(repo, {**base, "sector": ""}, "academic")
        assert again.id == finance.id
        db.commit()
        saved = contact_json(repo, repo.get(Contact, finance.id))
    assert saved["domains"] == {
        "finance": {"sector": "Asset Management"},
        "academic": {"sector": "Medical Imaging"},
    }


def test_persona_draft_sequence_domains_and_academic_mock_writing(client):
    persona = academic_persona(client)
    finance_persona = client.post(
        "/api/personas", json={"label": "Finance", "data": {"name": "Fin User"}}
    ).json()
    academic_contact = manual_contact(client, "academic", "Professor Example")
    finance_contact = manual_contact(client, "finance", "Finance Example")

    assert [item["id"] for item in client.get("/api/personas?domain=academic").json()] == [
        persona["id"]
    ]
    with Session() as db:
        revision = db.query(PersonaRevision).filter_by(persona_id=persona["id"]).one()
        assert revision.domain == "academic"

    updated = client.put(
        "/api/personas/" + persona["id"],
        json={
            "label": "Updated academic profile",
            "version": persona["version"],
            "data": {**persona["data"], "skills": "Python, PyTorch"},
        },
    )
    assert updated.status_code == 200 and updated.json()["domain"] == "academic"
    assert (
        client.put(
            "/api/personas/" + persona["id"],
            json={
                "domain": "finance",
                "label": "Wrong domain",
                "version": updated.json()["version"],
                "data": updated.json()["data"],
            },
        ).status_code
        == 422
    )

    draft_response = client.post(
        "/api/drafts",
        json={
            "contact_id": academic_contact["id"],
            "persona_id": persona["id"],
            "purpose": "Ask whether our research interests align",
        },
    )
    assert draft_response.status_code == 200, draft_response.text
    draft = draft_response.json()
    assert draft["domain"] == "academic"
    assert draft["starting_point"] == "PhD Inquiry"

    generated = client.post(
        f'/api/drafts/{draft["id"]}/generate',
        json={"action": "generate", "revision": draft["revision"]},
    )
    assert generated.status_code == 200, generated.text
    generated_text = generated.json()["subject"] + generated.json()["body_html"]
    assert "research" in generated_text.lower()
    assert "finance roles" not in generated_text.lower()
    draft = generated.json()

    omitted_domain = dict(draft)
    omitted_domain.pop("domain")
    omitted_domain["subject"] = "Updated research question"
    saved = client.put("/api/drafts/" + draft["id"], json=omitted_domain)
    assert saved.status_code == 200 and saved.json()["domain"] == "academic"
    assert (
        client.put(
            "/api/drafts/" + draft["id"],
            json={**saved.json(), "domain": "finance"},
        ).status_code
        == 422
    )

    for payload in (
        {
            "domain": "academic",
            "contact_id": finance_contact["id"],
            "persona_id": persona["id"],
        },
        {
            "domain": "academic",
            "contact_id": academic_contact["id"],
            "persona_id": finance_persona["id"],
        },
    ):
        assert client.post("/api/drafts", json=payload).status_code == 422

    copied = client.post(
        "/api/sequences",
        json={"name": "Academic follow-up", "draft_ids": [draft["id"]]},
    )
    assert copied.status_code == 201, copied.text
    assert copied.json()["steps"][0]["draft"]["domain"] == "academic"
    templated = client.post(
        "/api/sequences",
        json={
            "name": "Academic template",
            "template_id": "default-networking",
            "persona_id": persona["id"],
            "contact_id": academic_contact["id"],
        },
    )
    assert templated.status_code == 201, templated.text
    assert templated.json()["domain"] == "academic"
    assert all(step["draft"]["domain"] == "academic" for step in templated.json()["steps"])

    cross_domain_update = client.put(
        "/api/sequences/" + copied.json()["id"],
        json={
            "revision": copied.json()["revision"],
            "contact_id": finance_contact["id"],
            "persona_id": finance_persona["id"],
        },
    )
    assert cross_domain_update.status_code == 422
    unchanged = client.get("/api/sequences/" + copied.json()["id"]).json()
    assert unchanged["domain"] == "academic"
    assert unchanged["contact_id"] == academic_contact["id"]
    assert unchanged["persona_id"] == persona["id"]

    dual_domain_contact = client.put(
        "/api/contacts/" + academic_contact["id"],
        json={
            "domain": "finance",
            "name": academic_contact["name"],
            "title": academic_contact["title"],
            "company": academic_contact["company"],
            "sector": "Investment Banking",
        },
    )
    assert dual_domain_contact.status_code == 200
    assert set(dual_domain_contact.json()["domains"]) == {"finance", "academic"}
    finance_draft = client.post(
        "/api/drafts",
        json={
            "domain": "finance",
            "contact_id": academic_contact["id"],
            "persona_id": finance_persona["id"],
        },
    ).json()
    mixed = client.post(
        "/api/sequences",
        json={
            "name": "Mixed domains",
            "draft_ids": [draft["id"], finance_draft["id"]],
        },
    )
    assert mixed.status_code == 422

    finance_contact_as_academic = client.put(
        "/api/contacts/" + finance_contact["id"],
        json={
            "domain": "academic",
            "name": finance_contact["name"],
            "title": "Managing Director",
            "company": "Aster Capital",
            "sector": "Medical imaging",
        },
    ).json()
    factual_draft = client.post(
        "/api/drafts",
        json={
            "domain": "academic",
            "contact_id": finance_contact_as_academic["id"],
            "persona_id": persona["id"],
            "purpose": "Ask about research fit",
        },
    ).json()
    factual_generation = client.post(
        f'/api/drafts/{factual_draft["id"]}/generate',
        json={"action": "generate", "revision": factual_draft["revision"]},
    )
    assert factual_generation.status_code == 200
    assert "your research at aster capital" not in factual_generation.json()[
        "body_html"
    ].lower()
    assert "medical imaging" in factual_generation.json()["body_html"].lower()

    assert (
        client.post(
            "/api/drafts",
            json={
                "domain": "finance",
                "contact_id": finance_contact["id"],
                "persona_id": finance_persona["id"],
                "starting_point": "PhD Inquiry",
            },
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/drafts",
            json={
                "domain": "academic",
                "contact_id": academic_contact["id"],
                "persona_id": persona["id"],
                "starting_point": "Networking",
            },
        ).status_code
        == 422
    )


def academic_resume_bytes():
    document = Document()
    document.add_paragraph("Taylor Scholar")
    document.add_paragraph("Research Interests")
    document.add_paragraph("Multimodal medical imaging")
    document.add_paragraph("Research Experience")
    document.add_paragraph("Built a reproducible imaging pipeline")
    document.add_paragraph("Academic Goals")
    document.add_paragraph("Pursue doctoral research")
    document.add_paragraph("Target Institutions")
    document.add_paragraph("Example University")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_document_domain_cache_parse_and_persona_attachment(paused_workers):
    client = paused_workers
    content = academic_resume_bytes()

    def upload(domain):
        response = client.post(
            "/api/documents/jobs",
            data={"domain": domain},
            files={"file": ("cv.docx", content)},
        )
        assert response.status_code == 202, response.text
        return response.json()

    finance = upload("finance")
    academic = upload("academic")
    assert finance["id"] != academic["id"]
    assert upload("academic")["id"] == academic["id"]
    assert finance["domain"] == "finance" and academic["domain"] == "academic"
    assert process_document(academic["id"])
    parsed = client.get(f'/api/documents/{academic["id"]}/status').json()
    assert parsed["data"]["sectors"] == "Multimodal medical imaging"
    assert "imaging pipeline" in parsed["data"]["experience"]
    assert parsed["data"]["career_goals"] == "Pursue doctoral research"
    assert parsed["data"]["target_roles"] == "Example University"
    assert [item["id"] for item in client.get("/api/documents?domain=academic").json()] == [
        academic["id"]
    ]

    attached = client.post(
        "/api/personas",
        json={
            "domain": "academic",
            "label": "Parsed academic profile",
            "document_id": academic["id"],
            "data": parsed["data"],
        },
    )
    assert attached.status_code == 200, attached.text
    mismatch = client.post(
        "/api/personas",
        json={
            "label": "Wrong document domain",
            "document_id": academic["id"],
            "data": parsed["data"],
        },
    )
    assert mismatch.status_code == 422
