import re
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from pypdf import PdfReader
from docx import Document
from ..schemas import PersonaData

HEADINGS = {
    "education": "education",
    "教育经历": "education",
    "教育背景": "education",
    "experience": "experience",
    "work experience": "experience",
    "工作经历": "experience",
    "professional experience": "experience",
    "skills": "skills",
    "技能": "skills",
    "sectors": "sectors",
    "金融领域": "sectors",
    "career goals": "career_goals",
    "职业目标": "career_goals",
    "target regions": "target_regions",
    "目标地区": "target_regions",
    "target roles": "target_roles",
    "目标机构或职位": "target_roles",
    "contact purpose": "contact_purpose",
    "联系目的": "contact_purpose",
}


def extract_text(content: bytes, extension: str):
    if extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise ValueError("This file is not a valid PDF. Use manual entry instead.")
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise ValueError(
                "Password-protected PDF is unsupported. Upload an unlocked copy or enter your background manually."
            )
        if len(reader.pages) > 30:
            raise ValueError("Resume exceeds the 30-page limit. Upload a shorter file.")
        text = "\n".join(p.extract_text() or "" for p in reader.pages)
    elif extension == ".docx":
        try:
            with ZipFile(BytesIO(content)) as z:
                if (
                    sum(i.file_size for i in z.infolist()) > 25 * 1024 * 1024
                    or len(z.infolist()) > 1000
                ):
                    raise ValueError("DOCX expands beyond the safe size limit.")
                if "word/document.xml" not in z.namelist():
                    raise ValueError("Invalid DOCX document.")
            doc = Document(BytesIO(content))
            text = "\n".join(
                [p.text for p in doc.paragraphs]
                + [c.text for t in doc.tables for row in t.rows for c in row.cells]
            )
        except BadZipFile:
            raise ValueError("This file is not a valid DOCX. Use manual entry instead.")
    else:
        raise ValueError("Only text-based PDF and DOCX resumes are supported.")
    if len(re.sub(r"\s", "", text)) < 30:
        raise ValueError(
            "No usable text was found. Scanned/image-only resumes are unsupported (no OCR). Upload a text-based PDF/DOCX or use manual entry."
        )
    if len(text) > 50000:
        raise ValueError("Extracted text is too long (maximum 50,000 characters).")
    return text


def extract_sections(text):
    data = PersonaData().model_dump()
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    first = lines[0] if lines else ""
    if 1 < len(first) < 80 and not any(
        x in first.lower() for x in ("@", "resume", "curriculum", "简历", "http")
    ):
        data["name"] = first
    field = None
    for line in lines[1:]:
        parts = re.split(r"[:：]", line, maxsplit=1)
        key = HEADINGS.get(parts[0].lower().strip())
        if key:
            field = key
            if len(parts) > 1:
                data[field] += parts[1].strip() + "\n"
        elif field:
            data[field] += line + "\n"
    return {k: v.strip() for k, v in data.items()}
